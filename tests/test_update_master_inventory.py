import sys
from pathlib import Path
from typing import cast
from unittest.mock import MagicMock, patch

import openpyxl
import pytest
from openpyxl import Workbook
from openpyxl.cell.cell import Cell, MergedCell
from openpyxl.worksheet.worksheet import Worksheet

from update_master_inventory import (
    coerce_for_cell,
    column_letter,
    count_master_data_rows,
    find_header_row_number,
    get_active_worksheet,
    get_cell,
    main,
    normalize_value,
    update_xlsx,
)


class TestNormalization:
    """Valida conversiones de celdas, formato de columnas y coerción numéricas."""

    @pytest.mark.parametrize(
        "value, expected",
        [
            (None, ""),
            ("16.0", "16"),
            ("16", "16"),
            ("16.5", "16.5"),
            (" text ", "text"),
            (16.0, "16"),
            (16.5, "16.5"),
            (42, "42"),
        ],
    )
    def test_normalize_value(self, value, expected) -> None:
        assert normalize_value(value) == expected

    @pytest.mark.parametrize(
        "value, expected",
        [
            ("", None),
            ("16", 16),
            ("-42", -42),
            ("16.5", 16.5),
            ("text", "text"),
            (" 42 ", 42),
        ],
    )
    def test_coerce_for_cell(self, value, expected) -> None:
        assert coerce_for_cell(value) == expected

    @pytest.mark.parametrize(
        "index, expected",
        [
            (0, "A"),
            (25, "Z"),
            (26, "AA"),
            (27, "AB"),
            (51, "AZ"),
            (52, "BA"),
            (701, "ZZ"),
        ],
    )
    def test_column_letter(self, index, expected) -> None:
        assert column_letter(index) == expected


class TestWorksheetValidation:
    """Valida la interacción con las estructuras de openpyxl."""

    def test_get_active_worksheet_success(self) -> None:
        wb = Workbook()
        ws = get_active_worksheet(wb)
        assert isinstance(ws, Worksheet)

    def test_get_active_worksheet_fails_on_chartsheet(self) -> None:
        wb = MagicMock(spec=Workbook)
        wb.active = MagicMock()  # Fake no-worksheet
        with pytest.raises(SystemExit):
            get_active_worksheet(wb)

    def test_get_cell_success(self) -> None:
        wb = Workbook()
        ws = cast(Worksheet, wb.active)
        cell = get_cell(ws, 1, 1)  # A1
        assert isinstance(cell, Cell)

    def test_get_cell_fails_on_merged_cell(self) -> None:
        wb = Workbook()
        ws = cast(Worksheet, wb.active)
        with (
            patch.object(ws, "cell", return_value=MergedCell(ws)),
            pytest.raises(SystemExit),
        ):
            get_cell(ws, 1, 1)

    def test_find_header_row_number(self) -> None:
        wb = Workbook()
        ws = cast(Worksheet, wb.active)
        ws.append(["Inventario", "General"])
        ws.append(["UUID", "IP", "MAC"])
        ws.append(["123", "10.0.0.1", "aa:bb"])

        merged_header = ["UUID", "IP", "MAC"]
        assert find_header_row_number(ws, merged_header) == 2

    def test_find_header_row_number_not_found(self) -> None:
        wb = Workbook()
        ws = cast(Worksheet, wb.active)
        ws.append(["UUID", "MAC"])  # Faltaría IP

        merged_header = ["UUID", "IP", "MAC"]
        assert find_header_row_number(ws, merged_header) is None

    def test_count_master_data_rows(self) -> None:
        wb = Workbook()
        ws = cast(Worksheet, wb.active)
        ws.append(["UUID", "IP"])  # Row 1 (header)
        ws.append(["123", "10.0.0.1"])  # Row 2 (data)
        ws.append(["456", "10.0.0.2"])  # Row 3 (data)
        ws.append([None, ""])  # Row 4 (empty)
        ws.append([None, None])  # Row 5 (empty)

        count = count_master_data_rows(ws, header_row_number=1, n_cols=2)
        assert count == 2


class TestUpdateXlsx:
    """Valida la detección de cambios, mutación y salvaguardas."""

    def test_update_xlsx_happy_path(self, tmp_path: Path) -> None:
        # Preparar archivo
        wb = Workbook()
        ws = cast(Worksheet, wb.active)
        ws.append(["Inventario"])  # Fila 1 (basura)
        ws.append(["UUID", "IP"])  # Fila 2 (header)
        ws.append(["123", "10.0.0.1"])  # Fila 3 (data a modificar)
        ws.append(["456", "10.0.0.2"])  # Fila 4 (data intacta)

        xlsx_path = tmp_path / "master.xlsx"
        wb.save(xlsx_path)

        merged_header = ["UUID", "IP"]
        # IP cambió a 10.0.0.99
        merged_rows = [["123", "10.0.0.99"], ["456", "10.0.0.2"]]

        changes = update_xlsx(xlsx_path, merged_header, merged_rows)

        assert changes == {"IP": ["B3"]}

        # Validar el archivo editado
        edited_wb = openpyxl.load_workbook(xlsx_path)
        edited_ws = cast(Worksheet, edited_wb.active)
        assert edited_ws["A3"].value == "123"  # No cambió, preserva su valor original
        assert edited_ws["B3"].value == "10.0.0.99"
        assert edited_ws["A4"].value == "456"
        assert edited_ws["B4"].value == "10.0.0.2"

    def test_update_xlsx_header_not_found(self, tmp_path: Path) -> None:
        wb = Workbook()
        ws = cast(Worksheet, wb.active)
        ws.append(["A", "B"])
        xlsx_path = tmp_path / "master.xlsx"
        wb.save(xlsx_path)

        with pytest.raises(SystemExit):
            update_xlsx(xlsx_path, ["UUID"], [["123"]])

    def test_update_xlsx_row_mismatch(self, tmp_path: Path) -> None:
        wb = Workbook()
        ws = cast(Worksheet, wb.active)
        ws.append(["UUID", "IP"])
        ws.append(["123", "10.0.0.1"])
        xlsx_path = tmp_path / "master.xlsx"
        wb.save(xlsx_path)

        # merged_rows tiene 2 filas de datos, el master tiene 1
        merged_rows = [["123", "10.0.0.1"], ["456", "10.0.0.2"]]

        with pytest.raises(SystemExit):
            update_xlsx(xlsx_path, ["UUID", "IP"], merged_rows)


class TestMainUpdateMasterInventory:
    """Valida la integración principal (IO, Extensiones, Cleanups)."""

    @pytest.fixture
    def setup_files(self, tmp_path: Path) -> tuple[Path, Path, Path]:
        merged_csv = tmp_path / "merged_inventory.csv"
        master_xlsx = tmp_path / "master_inventory.xlsx"
        output_xlsx = tmp_path / "master_inventory_updated.xlsx"

        merged_csv.touch()
        master_xlsx.touch()
        return merged_csv, master_xlsx, output_xlsx

    @patch("update_master_inventory.load_csv_data")
    @patch("update_master_inventory.validate_rows")
    @patch("update_master_inventory.update_xlsx", return_value={})
    def test_main_happy_path_xlsx(
        self,
        mock_update: MagicMock,
        mock_val: MagicMock,
        mock_load: MagicMock,
        setup_files: tuple[Path, Path, Path],
    ) -> None:
        merged_csv, master_xlsx, output_xlsx = setup_files

        mock_load.return_value = (
            '"UUID","IP"\n',
            ['"UUID"', '"IP"'],
            [['"123"', '"10.0.0.1"']],
        )
        mock_val.return_value = [['"123"', '"10.0.0.1"']]

        with patch.object(
            sys,
            "argv",
            ["script.py", str(merged_csv), str(master_xlsx), str(output_xlsx)],
        ):
            main()

        assert output_xlsx.exists()
        mock_update.assert_called_once()
        args, _kwargs = mock_update.call_args
        assert args[1] == ["UUID", "IP"]
        assert args[2] == [["123", "10.0.0.1"]]

    def test_main_missing_files(self, setup_files: tuple[Path, Path, Path]) -> None:
        merged_csv, master_xlsx, output_xlsx = setup_files
        merged_csv.unlink()  # Delete merged_csv

        with (
            patch.object(
                sys,
                "argv",
                ["script.py", str(merged_csv), str(master_xlsx), str(output_xlsx)],
            ),
            pytest.raises(SystemExit),
        ):
            main()

    def test_main_unsupported_extension(
        self, setup_files: tuple[Path, Path, Path]
    ) -> None:
        merged_csv, master_xlsx, _ = setup_files
        # Cambiamos extensión a .txt
        master_txt = master_xlsx.with_suffix(".txt")
        master_txt.touch()

        with (
            patch.object(sys, "argv", ["script.py", str(merged_csv), str(master_txt)]),
            pytest.raises(SystemExit),
        ):
            main()

    def test_main_extension_mismatch(
        self, setup_files: tuple[Path, Path, Path]
    ) -> None:
        merged_csv, master_xlsx, _ = setup_files
        # Salida es ods pero master es xlsx
        output_ods = master_xlsx.parent / "master_updated.ods"

        with (
            patch.object(
                sys,
                "argv",
                ["script.py", str(merged_csv), str(master_xlsx), str(output_ods)],
            ),
            pytest.raises(SystemExit),
        ):
            main()

    def test_main_input_output_collision(
        self, setup_files: tuple[Path, Path, Path]
    ) -> None:
        merged_csv, master_xlsx, _ = setup_files
        # master == output
        with (
            patch.object(
                sys,
                "argv",
                ["script.py", str(merged_csv), str(master_xlsx), str(master_xlsx)],
            ),
            pytest.raises(SystemExit),
        ):
            main()

    @patch("update_master_inventory.load_csv_data")
    @patch("update_master_inventory.validate_rows")
    @patch("update_master_inventory.update_xlsx")
    def test_main_temp_file_cleanup_on_failure(
        self,
        mock_update: MagicMock,
        mock_val: MagicMock,
        mock_load: MagicMock,
        setup_files: tuple[Path, Path, Path],
    ) -> None:
        merged_csv, master_xlsx, output_xlsx = setup_files

        mock_load.return_value = ('"UUID"\n', ['"UUID"'], [['"123"']])
        mock_val.return_value = [['"123"']]
        # Forzamos fallo (SystemExit simulando un update_xlsx abortando)
        mock_update.side_effect = SystemExit()

        with (
            patch.object(
                sys,
                "argv",
                ["script.py", str(merged_csv), str(master_xlsx), str(output_xlsx)],
            ),
            pytest.raises(SystemExit),
        ):
            main()

        assert not output_xlsx.exists()
        # Verificar que no quedan archivos *.tmp
        tmp_files = list(output_xlsx.parent.glob(f".{output_xlsx.name}.*.tmp"))
        assert len(tmp_files) == 0
